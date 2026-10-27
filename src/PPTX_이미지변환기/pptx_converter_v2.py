# Ranet Toolkit - Copyright (C) 2026 Region & Agriculture NETwork Cooperative, Jeonnam
# SPDX-License-Identifier: AGPL-3.0-only  (see LICENSE)
import os
import sys
import shutil
import tempfile
import winreg
import zipfile
import io
import ctypes

from PySide6.QtWidgets import (
    QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout,
    QLabel, QPushButton, QTextEdit, QFrame, QProgressBar,
    QFileDialog, QMessageBox, QSizePolicy, QScrollArea, QStackedWidget,
)
from PySide6.QtCore import Qt, Signal, QObject, QThread, Slot, QRectF
from PySide6.QtGui import (
    QFont, QFontDatabase, QPixmap, QIcon,
    QDragEnterEvent, QDropEvent, QPalette, QColor,
    QPainter, QPainterPath,
)

import pythoncom
import win32com.client
from PIL import Image
from pptx import Presentation


# ── 리소스 경로 ──────────────────────────────────
def _resource(filename: str) -> str:
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, filename)


# ── 색상 상수 ────────────────────────────────────
BG      = "#F5F6FA"
SURFACE = "#FFFFFF"
ACCENT  = "#2B7FFF"
DANGER  = "#EF4444"
SUCCESS = "#22C55E"
BORDER  = "#E8ECF2"
TXT     = "#1A1D23"
TXT_S   = "#6B7280"
TXT_T   = "#9BA3AF"
PRV_BG  = "#F0F2F7"

# ── 버튼 인라인 스타일 (QSS 선택자 오작동 방지) ──
_BTN_DPI_ON   = f"background:{ACCENT};color:white;border:none;border-radius:6px;"
_BTN_DPI_OFF  = f"background:{BORDER};color:{TXT_S};border:none;border-radius:6px;"
_BTN_START_ON  = f"background:{ACCENT};color:white;border:none;border-radius:8px;"
_BTN_START_OFF = "background:#D1D5DB;color:#9CA3AF;border:none;border-radius:8px;"
_BTN_STOP_IDLE   = f"background:{BORDER};color:{TXT_T};border:none;border-radius:8px;"
_BTN_STOP_ACTIVE = f"background:{DANGER};color:white;border:none;border-radius:8px;"
_BTN_BROWSE_SS   = f"background:{ACCENT};color:white;border:none;border-radius:8px;padding:0 20px;"
_BTN_RESEL_SS    = f"background:{BORDER};color:{TXT_S};border:none;border-radius:5px;padding:0 10px;"
_BTN_RESEL_DIS   = f"background:{BORDER};color:{TXT_T};border:none;border-radius:5px;padding:0 10px;"


# ── 폰트 ─────────────────────────────────────────
_PF: dict[str, str] = {}


def _load_fonts() -> None:
    d = _resource("fonts")
    for wt, fn in (
        ("regular",  "Pretendard-Regular.ttf"),
        ("medium",   "Pretendard-Medium.ttf"),
        ("semibold", "Pretendard-SemiBold.ttf"),
        ("light",    "Pretendard-Light.ttf"),
    ):
        p = os.path.join(d, fn)
        if os.path.exists(p):
            fid = QFontDatabase.addApplicationFont(p)
            if fid >= 0:
                fams = QFontDatabase.applicationFontFamilies(fid)
                if fams:
                    _PF[wt] = fams[0]


def _font(pt: int, wt: str = "regular") -> QFont:
    fam = _PF.get(wt) or _PF.get("regular") or "맑은 고딕"
    w_map = {
        "semibold": QFont.Weight.DemiBold,
        "bold":     QFont.Weight.Bold,
        "medium":   QFont.Weight.Medium,
        "regular":  QFont.Weight.Normal,
        "light":    QFont.Weight.Light,
    }
    f = QFont(fam, pt)
    f.setWeight(w_map.get(wt, QFont.Weight.Normal))
    return f


# ── 라이트 팔레트 (QMessageBox 등 시스템 다이얼로그 흰색 테마) ──
def _make_light_palette() -> QPalette:
    p = QPalette()
    pairs = [
        (QPalette.ColorRole.Window,          QColor(BG)),
        (QPalette.ColorRole.WindowText,      QColor(TXT)),
        (QPalette.ColorRole.Base,            QColor(SURFACE)),
        (QPalette.ColorRole.AlternateBase,   QColor(BG)),
        (QPalette.ColorRole.Text,            QColor(TXT)),
        (QPalette.ColorRole.Button,          QColor(BORDER)),
        (QPalette.ColorRole.ButtonText,      QColor(TXT)),
        (QPalette.ColorRole.BrightText,      QColor(SURFACE)),
        (QPalette.ColorRole.Link,            QColor(ACCENT)),
        (QPalette.ColorRole.Highlight,       QColor(ACCENT)),
        (QPalette.ColorRole.HighlightedText, QColor(SURFACE)),
        (QPalette.ColorRole.ToolTipBase,     QColor(SURFACE)),
        (QPalette.ColorRole.ToolTipText,     QColor(TXT)),
    ]
    for role, color in pairs:
        p.setColor(role, color)
        p.setColor(QPalette.ColorGroup.Disabled, role, color.lighter(120))
    return p


# ── 레지스트리 헬퍼 ──────────────────────────────
_PPT_VERS = ["16.0", "15.0", "14.0", "12.0", "11.0"]
_REG_SUB  = r"Software\Microsoft\Office\{v}\PowerPoint\Options"
_REG_VAL  = "ExportBitmapResolution"


def _reg_find() -> str | None:
    for v in _PPT_VERS:
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, _REG_SUB.format(v=v)):
                return _REG_SUB.format(v=v)
        except OSError:
            pass
    return None


def _reg_get(sub: str, name: str):
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, sub) as k:
            v, _ = winreg.QueryValueEx(k, name)
            return v
    except OSError:
        return None


def _reg_set(sub: str, name: str, val: int) -> None:
    with winreg.OpenKey(winreg.HKEY_CURRENT_USER, sub,
                        access=winreg.KEY_SET_VALUE) as k:
        winreg.SetValueEx(k, name, 0, winreg.REG_DWORD, val)


def _reg_del(sub: str, name: str) -> None:
    try:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, sub,
                            access=winreg.KEY_SET_VALUE) as k:
            winreg.DeleteValue(k, name)
    except OSError:
        pass


class _Cancelled(Exception):
    pass


# ── 변환 워커 ────────────────────────────────────
class ConvertWorker(QObject):
    log       = Signal(str)
    progress  = Signal(int, int, str)
    done      = Signal(str)
    error     = Signal(str)
    cancelled = Signal()

    def __init__(self, input_path: str, dpi: int) -> None:
        super().__init__()
        self._path  = input_path
        self._dpi   = dpi
        self._stop  = False
        self._ppt   = None
        self._owned = False
        self._out   = None
        self._tmp   = None

    def request_stop(self) -> None:
        self._stop = True

    @Slot()
    def run(self) -> None:
        pythoncom.CoInitialize()
        self._tmp = tempfile.mkdtemp(prefix="pptx_conv_")
        safe = os.path.join(self._tmp, "source.pptx")
        reg = _reg_find()
        orig = None
        aborted = False
        try:
            self.log.emit("=== 변환 시작 ===")
            if reg:
                orig = _reg_get(reg, _REG_VAL)
                _reg_set(reg, _REG_VAL, self._dpi)
                self.log.emit(f"DPI 설정: {self._dpi}dpi")
            else:
                self.log.emit("PowerPoint 레지스트리 없음 — DPI 설정 건너뜀")
            shutil.copy2(self._path, safe)
            self.log.emit("[1단계] PNG 내보내기 중...")
            pngs = self._export_png(safe)
            self.log.emit("[2단계] 슬라이드에 이미지 삽입 중...")
            out = self._insert(safe, pngs)
            self.log.emit("=== 완료 ===")
            self.done.emit(out)
        except _Cancelled:
            aborted = True
            self.log.emit("=== 변환 중지됨 ===")
            self.cancelled.emit()
        except Exception as ex:
            self.log.emit(f"[오류] {ex}")
            self.error.emit(str(ex))
        finally:
            if reg:
                if orig is None:
                    _reg_del(reg, _REG_VAL)
                else:
                    _reg_set(reg, _REG_VAL, orig)
            if aborted and self._out and os.path.exists(self._out):
                try:
                    os.remove(self._out)
                except Exception:
                    pass
            tmp, self._tmp = self._tmp, None
            shutil.rmtree(tmp, ignore_errors=True)
            self.log.emit("(임시 파일 삭제 완료)")
            pythoncom.CoUninitialize()

    def _export_png(self, safe: str) -> list:
        prs_com = None
        pngs = []
        try:
            try:
                self._ppt = win32com.client.GetActiveObject("PowerPoint.Application")
                self._owned = False
            except Exception:
                self._ppt = win32com.client.Dispatch("PowerPoint.Application")
                self._owned = True
            prs_com = self._ppt.Presentations.Open(
                safe, ReadOnly=True, Untitled=False, WithWindow=False)
            n = prs_com.Slides.Count
            w = int(prs_com.PageSetup.SlideWidth  / 72 * self._dpi)
            h = int(prs_com.PageSetup.SlideHeight / 72 * self._dpi)
            self.log.emit(f"슬라이드 크기: {w}x{h}px | 총 {n}슬라이드")
            for i in range(1, n + 1):
                if self._stop:
                    raise _Cancelled()
                png = os.path.join(self._tmp, f"slide_{i:04d}.png")
                prs_com.Slides(i).Export(png, "PNG", w, h)
                pngs.append(png)
                self.progress.emit(i, n, f"[1단계] PNG 내보내기: {i}/{n} 슬라이드")
        finally:
            if prs_com:
                prs_com.Close()
            if self._ppt and self._owned:
                try:
                    self._ppt.Quit()
                except Exception:
                    pass
            self._ppt = None
            self._owned = False
        return pngs

    def _insert(self, safe: str, pngs: list) -> str:
        orig = Presentation(safe)
        prs = Presentation()
        prs.slide_width  = orig.slide_width
        prs.slide_height = orig.slide_height
        blank = prs.slide_layouts[6]
        n = len(pngs)
        for i, png in enumerate(pngs):
            if self._stop:
                raise _Cancelled()
            sl = prs.slides.add_slide(blank)
            sl.shapes.add_picture(png, 0, 0, prs.slide_width, prs.slide_height)
            self.progress.emit(i + 1, n, f"[2단계] 이미지 삽입: {i+1}/{n} 슬라이드")
        base, ext = os.path.splitext(self._path)
        self._out = base + "_이미지변환" + ext
        prs.save(self._out)
        self.log.emit(f"저장: {os.path.basename(self._out)}")
        return self._out


# ── 텍스트 말줄임 레이블 ─────────────────────────
class ElideLabel(QLabel):
    def __init__(self, text: str = "", parent=None) -> None:
        super().__init__(parent)
        self._full = text
        self.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Preferred)
        self.setMinimumWidth(0)

    def setText(self, text: str) -> None:
        self._full = text
        self._apply()

    def resizeEvent(self, event) -> None:
        super().resizeEvent(event)
        self._apply()

    def _apply(self) -> None:
        fm = self.fontMetrics()
        elided = fm.elidedText(self._full, Qt.TextElideMode.ElideMiddle,
                                max(self.width() - 4, 0))
        super().setText(elided)


# ── 둥근 클리핑 미리보기 위젯 ───────────────────────
class RoundedPreviewWidget(QWidget):
    """카드의 border-radius에 맞춰 썸네일을 QPainterPath로 클리핑해 렌더링."""
    RADIUS = 10  # card border-radius(12) - border(1.5) ≈ 10

    def __init__(self, parent=None) -> None:
        super().__init__(parent)
        self._px: QPixmap | None = None
        self.setAttribute(Qt.WidgetAttribute.WA_OpaquePaintEvent)

    def setPixmap(self, px: QPixmap) -> None:
        self._px = px
        self.update()

    def clearPixmap(self) -> None:
        self._px = None
        self.update()

    def paintEvent(self, _) -> None:
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)

        clip = QPainterPath()
        clip.addRoundedRect(QRectF(self.rect()), self.RADIUS, self.RADIUS)
        painter.setClipPath(clip)

        painter.fillRect(self.rect(), QColor(PRV_BG))

        if self._px:
            w, h = self.width(), self.height()
            scaled = self._px.scaled(w, h,
                                     Qt.AspectRatioMode.KeepAspectRatio,
                                     Qt.TransformationMode.SmoothTransformation)
            x = (w - scaled.width())  // 2
            y = (h - scaled.height()) // 2
            painter.drawPixmap(x, y, scaled)
        else:
            painter.setPen(QColor(TXT_T))
            painter.setFont(_font(9, "light"))
            painter.drawText(self.rect(), Qt.AlignmentFlag.AlignCenter, "미리보기")


# ── QSS (구조·카드·공통 위젯) ──────────────────────
_QSS = f"""
* {{ outline: none; }}

QMainWindow, QWidget#root, QWidget#scroll_content {{
    background: {BG};
}}
QScrollArea {{ border: none; background: {BG}; }}

QFrame#card {{
    background: {SURFACE};
    border: 1px solid {BORDER};
    border-radius: 12px;
}}
QFrame#drop_card {{
    background: {SURFACE};
    border: 1.5px solid {ACCENT};
    border-radius: 12px;
}}

QProgressBar {{
    background: {BORDER};
    border: none;
    border-radius: 3px;
}}
QProgressBar::chunk {{
    background: {ACCENT};
    border-radius: 3px;
}}

QTextEdit#log_box {{
    background: #FAFBFC;
    border: 1px solid {BORDER};
    border-radius: 8px;
    color: {TXT};
    padding: 4px;
    selection-background-color: #DBEAFE;
}}

QScrollBar:vertical {{
    background: {BG};
    width: 8px;
    border-radius: 4px;
    margin: 0;
}}
QScrollBar::handle:vertical {{
    background: #D1D5DB;
    border-radius: 4px;
    min-height: 24px;
}}
QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical {{ height: 0; border: none; }}
QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {{ background: none; }}

QMessageBox {{
    background: {SURFACE};
}}
QMessageBox QLabel {{
    color: {TXT};
    background: transparent;
}}
QMessageBox QPushButton {{
    background: {ACCENT};
    color: white;
    border: none;
    border-radius: 6px;
    min-width: 80px;
    min-height: 30px;
    padding: 4px 16px;
}}
QMessageBox QPushButton:hover {{ background: #1A6FEF; }}
"""


# ── 메인 윈도우 ───────────────────────────────────
class App(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("PPTX 이미지 변환기 V4")
        self.setMinimumSize(480, 560)
        self.resize(540, 640)
        self.setAcceptDrops(True)
        self.setStyleSheet(_QSS)

        self._file_path   = ""
        self._dpi         = 300
        self._dpi_btns: dict[int, QPushButton] = {}
        self._worker: ConvertWorker | None = None
        self._thread: QThread | None       = None
        self._preview_px: QPixmap | None   = None

        self._apply_icon()
        self._build_ui()

    def _apply_icon(self) -> None:
        for name in ("logo.ico", "logo.png"):
            p = _resource(name)
            if os.path.exists(p):
                self.setWindowIcon(QIcon(p))
                break

    # ── UI 구성 ───────────────────────────────────
    def _build_ui(self) -> None:
        root = QWidget()
        root.setObjectName("root")
        self.setCentralWidget(root)

        vl = QVBoxLayout(root)
        vl.setContentsMargins(0, 0, 0, 0)
        vl.setSpacing(0)

        scroll = QScrollArea()
        scroll.setWidgetResizable(True)
        scroll.setHorizontalScrollBarPolicy(Qt.ScrollBarPolicy.ScrollBarAlwaysOff)

        sc = QWidget()
        sc.setObjectName("scroll_content")
        cl = QVBoxLayout(sc)
        cl.setContentsMargins(20, 16, 20, 0)
        cl.setSpacing(0)

        self._build_drop_zone(cl)
        self._build_dpi_card(cl)
        self._build_progress_card(cl)
        self._build_action_buttons(cl)
        self._build_log_card(cl)

        footer = QLabel("RANET JN")
        footer.setFont(_font(8, "light"))
        footer.setStyleSheet(f"color: {TXT_T}; background: transparent;")
        footer.setAlignment(Qt.AlignmentFlag.AlignCenter)
        footer.setContentsMargins(0, 2, 0, 8)
        cl.addWidget(footer)

        scroll.setWidget(sc)
        vl.addWidget(scroll)

    # ── 드롭존 (파일 선택 전: 드롭 UI / 선택 후: 썸네일) ──
    def _build_drop_zone(self, layout: QVBoxLayout) -> None:
        self._drop_card = QFrame()
        self._drop_card.setObjectName("drop_card")
        self._drop_card.setFixedHeight(178)

        card_l = QVBoxLayout(self._drop_card)
        card_l.setContentsMargins(0, 0, 0, 0)
        card_l.setSpacing(0)

        self._drop_stack = QStackedWidget()
        card_l.addWidget(self._drop_stack)

        # ── 페이지 0: 빈 상태 ──
        empty_w = QWidget()
        empty_w.setStyleSheet("background: transparent;")
        el = QVBoxLayout(empty_w)
        el.setAlignment(Qt.AlignmentFlag.AlignCenter)
        el.setSpacing(6)
        el.setContentsMargins(20, 10, 20, 10)

        self._drop_icon = QLabel("📁")
        self._drop_icon.setFont(QFont("Segoe UI Emoji", 26))
        self._drop_icon.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self._drop_icon.setStyleSheet("background: transparent;")
        el.addWidget(self._drop_icon)

        title_lbl = QLabel("여기에 PPTX 파일을 드래그 앤 드롭")
        title_lbl.setFont(_font(10, "medium"))
        title_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        title_lbl.setStyleSheet(f"color: {TXT}; background: transparent;")
        el.addWidget(title_lbl)

        sub_lbl = QLabel("또는 아래 버튼으로 파일을 선택하세요")
        sub_lbl.setFont(_font(9, "regular"))
        sub_lbl.setAlignment(Qt.AlignmentFlag.AlignCenter)
        sub_lbl.setStyleSheet(f"color: {TXT_S}; background: transparent;")
        el.addWidget(sub_lbl)

        self._btn_browse = QPushButton("파일 선택하기")
        self._btn_browse.setFont(_font(9, "medium"))
        self._btn_browse.setCursor(Qt.CursorShape.PointingHandCursor)
        self._btn_browse.setFixedHeight(34)
        self._btn_browse.setStyleSheet(_BTN_BROWSE_SS)
        self._btn_browse.clicked.connect(self._browse)
        el.addWidget(self._btn_browse, alignment=Qt.AlignmentFlag.AlignCenter)

        self._drop_stack.addWidget(empty_w)

        # ── 페이지 1: 썸네일 미리보기 (둥근 클리핑 커스텀 위젯) ──
        self._prv_widget = RoundedPreviewWidget()
        self._drop_stack.addWidget(self._prv_widget)
        self._drop_stack.setCurrentIndex(0)

        layout.addWidget(self._drop_card)

        # ── 파일 정보 행 (파일 선택 후 카드 아래에 표시) ──
        self._file_info_row = QWidget()
        self._file_info_row.setVisible(False)
        self._file_info_row.setStyleSheet("background: transparent;")
        fil = QHBoxLayout(self._file_info_row)
        fil.setContentsMargins(4, 5, 4, 6)
        fil.setSpacing(8)

        chk = QLabel("✓")
        chk.setFont(_font(10, "semibold"))
        chk.setStyleSheet(f"color: {SUCCESS}; background: transparent;")
        fil.addWidget(chk)

        self._lbl_selected_name = ElideLabel("")
        self._lbl_selected_name.setFont(_font(9, "medium"))
        self._lbl_selected_name.setStyleSheet(f"color: {ACCENT}; background: transparent;")
        fil.addWidget(self._lbl_selected_name)

        self._btn_reselect = QPushButton("다른 파일 선택")
        self._btn_reselect.setFont(_font(8, "regular"))
        self._btn_reselect.setCursor(Qt.CursorShape.PointingHandCursor)
        self._btn_reselect.setFixedHeight(26)
        self._btn_reselect.setStyleSheet(_BTN_RESEL_SS)
        self._btn_reselect.clicked.connect(self._browse)
        fil.addWidget(self._btn_reselect, alignment=Qt.AlignmentFlag.AlignRight)

        layout.addWidget(self._file_info_row)
        layout.addSpacing(4)

    def _build_dpi_card(self, layout: QVBoxLayout) -> None:
        card = QFrame()
        card.setObjectName("card")
        vl = QVBoxLayout(card)
        vl.setContentsMargins(0, 0, 0, 0)
        vl.setSpacing(0)

        lbl = QLabel("내보내기 해상도 (DPI)")
        lbl.setFont(_font(9, "medium"))
        lbl.setStyleSheet(f"color: {TXT}; background: transparent;")
        lbl.setContentsMargins(14, 10, 14, 6)
        vl.addWidget(lbl)

        row_w = QWidget()
        row_w.setStyleSheet("background: transparent;")
        row_l = QHBoxLayout(row_w)
        row_l.setContentsMargins(14, 0, 14, 14)
        row_l.setSpacing(6)

        for dpi in (96, 150, 200, 300, 600, 900):
            btn = QPushButton(str(dpi))
            btn.setFont(_font(9, "medium"))
            btn.setCursor(Qt.CursorShape.PointingHandCursor)
            btn.setFixedHeight(32)
            btn.setSizePolicy(QSizePolicy.Policy.Expanding, QSizePolicy.Policy.Fixed)
            btn.setStyleSheet(_BTN_DPI_ON if dpi == self._dpi else _BTN_DPI_OFF)
            btn.clicked.connect(lambda _, d=dpi: self._select_dpi(d))
            row_l.addWidget(btn)
            self._dpi_btns[dpi] = btn

        vl.addWidget(row_w)
        layout.addWidget(card)
        layout.addSpacing(10)

    def _build_progress_card(self, layout: QVBoxLayout) -> None:
        card = QFrame()
        card.setObjectName("card")
        vl = QVBoxLayout(card)
        vl.setContentsMargins(14, 14, 14, 14)
        vl.setSpacing(0)

        top = QWidget()
        top.setStyleSheet("background: transparent;")
        tl = QHBoxLayout(top)
        tl.setContentsMargins(0, 0, 0, 0)

        prog_lbl = QLabel("진행 상황")
        prog_lbl.setFont(_font(9, "medium"))
        prog_lbl.setStyleSheet(f"color: {TXT}; background: transparent;")
        tl.addWidget(prog_lbl)
        tl.addStretch()

        self._lbl_count = QLabel("")
        self._lbl_count.setFont(_font(9, "regular"))
        self._lbl_count.setStyleSheet(f"color: {TXT_S}; background: transparent;")
        tl.addWidget(self._lbl_count)

        vl.addWidget(top)
        vl.addSpacing(8)

        self._progress = QProgressBar()
        self._progress.setFixedHeight(6)
        self._progress.setTextVisible(False)
        self._progress.setRange(0, 100)
        self._progress.setValue(0)
        vl.addWidget(self._progress)
        vl.addSpacing(8)

        self._lbl_status = QLabel("파일을 선택해 주세요.")
        self._lbl_status.setFont(_font(9, "regular"))
        self._lbl_status.setStyleSheet(f"color: {TXT_S}; background: transparent;")
        vl.addWidget(self._lbl_status)

        layout.addWidget(card)
        layout.addSpacing(10)

    def _build_action_buttons(self, layout: QVBoxLayout) -> None:
        row = QWidget()
        row.setStyleSheet("background: transparent;")
        rl = QHBoxLayout(row)
        rl.setContentsMargins(0, 0, 0, 10)
        rl.setSpacing(10)

        self._btn_start = QPushButton("▶  변환 시작")
        self._btn_start.setFont(_font(10, "medium"))
        self._btn_start.setCursor(Qt.CursorShape.PointingHandCursor)
        self._btn_start.setFixedHeight(44)
        self._btn_start.setEnabled(False)
        self._btn_start.setStyleSheet(_BTN_START_OFF)
        self._btn_start.clicked.connect(self._start)
        rl.addWidget(self._btn_start)

        self._btn_stop = QPushButton("■  변환 중지")
        self._btn_stop.setFont(_font(10, "medium"))
        self._btn_stop.setCursor(Qt.CursorShape.PointingHandCursor)
        self._btn_stop.setFixedHeight(44)
        self._btn_stop.setEnabled(False)
        self._btn_stop.setStyleSheet(_BTN_STOP_IDLE)
        self._btn_stop.clicked.connect(self._stop)
        rl.addWidget(self._btn_stop)

        layout.addWidget(row)

    def _build_log_card(self, layout: QVBoxLayout) -> None:
        card = QFrame()
        card.setObjectName("card")
        vl = QVBoxLayout(card)
        vl.setContentsMargins(0, 0, 0, 0)
        vl.setSpacing(0)

        lbl = QLabel("작업 로그")
        lbl.setFont(_font(9, "medium"))
        lbl.setStyleSheet(f"color: {TXT}; background: transparent;")
        lbl.setContentsMargins(14, 10, 14, 6)
        vl.addWidget(lbl)

        self._log_box = QTextEdit()
        self._log_box.setObjectName("log_box")
        self._log_box.setReadOnly(True)
        self._log_box.setFont(_font(9, "light"))
        self._log_box.setFixedHeight(120)
        wrap = QWidget()
        wrap.setStyleSheet("background: transparent;")
        wl = QVBoxLayout(wrap)
        wl.setContentsMargins(8, 0, 8, 8)
        wl.addWidget(self._log_box)
        vl.addWidget(wrap)

        layout.addWidget(card)
        layout.addSpacing(6)

    # ── 드래그앤드롭 ──────────────────────────────
    def dragEnterEvent(self, event: QDragEnterEvent) -> None:
        if event.mimeData().hasUrls():
            if any(u.toLocalFile().lower().endswith(".pptx")
                   for u in event.mimeData().urls()):
                event.acceptProposedAction()

    def dropEvent(self, event: QDropEvent) -> None:
        for url in event.mimeData().urls():
            p = url.toLocalFile()
            if p.lower().endswith(".pptx"):
                self._set_file(p)
                break

    # ── 파일 선택 ─────────────────────────────────
    def _browse(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "PPTX 파일 선택", "",
            "PowerPoint 파일 (*.pptx);;모든 파일 (*.*)")
        if path:
            self._set_file(path)

    def _set_file(self, path: str) -> None:
        self._file_path = path
        name = os.path.basename(path)

        # 드롭존 → 미리보기 페이지로 전환
        self._drop_stack.setCurrentIndex(1)
        self._file_info_row.setVisible(True)
        self._lbl_selected_name.setText(name)

        # 변환 시작 버튼 활성화
        self._btn_start.setEnabled(True)
        self._btn_start.setStyleSheet(_BTN_START_ON)

        self._progress.setValue(0)
        self._lbl_count.setText("")
        self._lbl_status.setText("변환 시작 버튼을 눌러 주세요.")
        self._lbl_status.setStyleSheet(f"color: {TXT_S}; background: transparent;")
        self._log_box.clear()
        self._log(f"선택됨: {name}")
        self._update_preview(path)

    def _update_preview(self, path: str) -> None:
        try:
            with zipfile.ZipFile(path, "r") as z:
                names = z.namelist()
                for cand in ("docProps/thumbnail.jpeg",
                             "docProps/thumbnail.jpg",
                             "docProps/thumbnail.png"):
                    if cand in names:
                        data = z.read(cand)
                        img = Image.open(io.BytesIO(data)).convert("RGB")
                        out_buf = io.BytesIO()
                        img.save(out_buf, format="PNG")
                        px = QPixmap()
                        px.loadFromData(out_buf.getvalue())
                        self._preview_px = px
                        # RoundedPreviewWidget이 paintEvent에서 스스로 스케일링·클리핑
                        self._prv_widget.setPixmap(px)
                        return
        except Exception:
            pass

    # ── DPI 선택 ──────────────────────────────────
    def _select_dpi(self, dpi: int) -> None:
        self._dpi = dpi
        for d, btn in self._dpi_btns.items():
            btn.setStyleSheet(_BTN_DPI_ON if d == dpi else _BTN_DPI_OFF)

    # ── 변환 시작/중지 ────────────────────────────
    def _start(self) -> None:
        if not self._file_path:
            return
        self._btn_start.setEnabled(False)
        self._btn_start.setStyleSheet(_BTN_START_OFF)
        self._btn_reselect.setEnabled(False)
        self._btn_reselect.setStyleSheet(_BTN_RESEL_DIS)
        self._btn_stop.setEnabled(True)
        self._btn_stop.setStyleSheet(_BTN_STOP_ACTIVE)
        self._progress.setValue(0)
        self._lbl_status.setText("변환 중...")
        self._lbl_status.setStyleSheet(f"color: {TXT}; background: transparent;")

        self._worker = ConvertWorker(self._file_path, self._dpi)
        self._thread = QThread()
        self._worker.moveToThread(self._thread)
        self._thread.started.connect(self._worker.run)
        self._worker.log.connect(self._log)
        self._worker.progress.connect(self._on_progress)
        self._worker.done.connect(self._on_done)
        self._worker.error.connect(self._on_error)
        self._worker.cancelled.connect(self._on_cancelled)
        self._worker.done.connect(self._thread.quit)
        self._worker.error.connect(self._thread.quit)
        self._worker.cancelled.connect(self._thread.quit)
        self._thread.finished.connect(self._on_finished)
        self._thread.start()

    def _stop(self) -> None:
        if self._worker:
            self._worker.request_stop()
        self._btn_stop.setEnabled(False)
        self._log("[중지] 현재 슬라이드 완료 후 중단됩니다.")
        self._lbl_status.setText("중지 중...")
        self._lbl_status.setStyleSheet(f"color: #F59E0B; background: transparent;")

    # ── 워커 시그널 ───────────────────────────────
    def _log(self, msg: str) -> None:
        self._log_box.append(msg)

    def _on_progress(self, current: int, total: int, status: str) -> None:
        pct = int(current / total * 100) if total else 0
        self._progress.setValue(pct)
        self._lbl_count.setText(f"{current} / {total}  ({pct}%)")
        self._lbl_status.setText(status)
        self._lbl_status.setStyleSheet(f"color: {TXT}; background: transparent;")

    def _on_done(self, out_path: str) -> None:
        self._progress.setValue(100)
        self._lbl_status.setText("변환 완료!")
        self._lbl_status.setStyleSheet(f"color: {SUCCESS}; background: transparent;")
        box = QMessageBox(self)
        box.setWindowTitle("완료")
        box.setText(f"변환이 완료되었습니다!\n\n저장 위치:\n{out_path}")
        box.setIcon(QMessageBox.Icon.Information)
        box.exec()

    def _on_error(self, msg: str) -> None:
        self._lbl_status.setText("오류 발생")
        self._lbl_status.setStyleSheet(f"color: {DANGER}; background: transparent;")
        box = QMessageBox(self)
        box.setWindowTitle("오류")
        box.setText(msg)
        box.setIcon(QMessageBox.Icon.Critical)
        box.exec()

    def _on_cancelled(self) -> None:
        self._progress.setValue(0)
        self._lbl_status.setText("변환 중지됨")
        self._lbl_status.setStyleSheet(f"color: #F59E0B; background: transparent;")

    def _on_finished(self) -> None:
        has_file = bool(self._file_path)
        self._btn_start.setEnabled(has_file)
        self._btn_start.setStyleSheet(_BTN_START_ON if has_file else _BTN_START_OFF)
        self._btn_reselect.setEnabled(True)
        self._btn_reselect.setStyleSheet(_BTN_RESEL_SS)
        self._btn_stop.setEnabled(False)
        self._btn_stop.setStyleSheet(_BTN_STOP_IDLE)
        self._lbl_count.setText("")

    # ── 타이틀바 흰색 (DWM) ──────────────────────
    def showEvent(self, event) -> None:
        super().showEvent(event)
        if not getattr(self, "_dwm_set", False):
            self._dwm_set = True
            try:
                hwnd = int(self.winId())
                ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, 20, ctypes.byref(ctypes.c_int(0)), 4)
                ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, 35, ctypes.byref(ctypes.c_int(0x00FFFFFF)), 4)
                ctypes.windll.dwmapi.DwmSetWindowAttribute(
                    hwnd, 36, ctypes.byref(ctypes.c_int(0x00000000)), 4)
            except Exception:
                pass

    # ── 종료 ──────────────────────────────────────
    def closeEvent(self, event) -> None:
        if self._thread and self._thread.isRunning():
            reply = QMessageBox.question(
                self, "종료 확인",
                "변환이 진행 중입니다.\n종료하시겠습니까?",
                QMessageBox.StandardButton.Yes | QMessageBox.StandardButton.No)
            if reply == QMessageBox.StandardButton.Yes:
                if self._worker:
                    self._worker.request_stop()
                self._thread.quit()
                self._thread.wait(3000)
                event.accept()
            else:
                event.ignore()
        else:
            event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    app.setPalette(_make_light_palette())
    _load_fonts()
    default_fam = _PF.get("regular", "맑은 고딕")
    app.setFont(QFont(default_fam, 10))
    window = App()
    window.show()
    sys.exit(app.exec())
